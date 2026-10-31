"""Presentation-only projection of existing VAT201 boxes for HTML and exports."""
import frappe

OUTPUT_KEYS = ["1" + chr(97 + i) for i in range(7)] + [str(i) for i in range(2, 9)]
INPUT_KEYS = ["9", "10", "11"]
NET_KEYS = ["12", "13", "14"]
ADJUSTMENT_KEYS = set(OUTPUT_KEYS[:7] + ["8", "9", "11"])


def form_data(result):
	boxes = {row["box"]: row for row in result["boxes"]}
	def row(key):
		box = boxes[key]
		fields = ["amount"] if key in ("4", "5") or key in NET_KEYS else ["amount", "vat_amount"]
		if key in ADJUSTMENT_KEYS:
			fields.append("adjustment")
		# Unexpected nonzero source values must remain visible, not silently masked.
		cells = [dict(field=field, value=box[field], applicable=field in fields or bool(box[field]),
			provisional=box.get("completeness", {}).get(field, {}).get("provisional", False))
			for field in ("amount", "vat_amount", "adjustment")]
		return dict(box=key, description=box["description"], cells=cells, total=key in ("8", "11", "14"))
	h = result["header"]
	return dict(header=h, status=result["status"], outputs=[row(key) for key in OUTPUT_KEYS],
		inputs=[row(key) for key in INPUT_KEYS], net=[row(key) for key in NET_KEYS],
		provisional=any(box.get("completeness", {}).get(field, {}).get("provisional")
			for box in boxes.values() for field in ("amount", "vat_amount", "adjustment")))


def form_html(result, interactive=False):
	return frappe.render_template("construction_management/report/uae_vat_201_return/form_content.html",
		dict(form=form_data(result), interactive=interactive))


def form_sheet(workbook, result):
	from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
	from openpyxl.comments import Comment
	from openpyxl.worksheet.pagebreak import Break
	form = form_data(result)
	h = form["header"]
	sheet = workbook.create_sheet("VAT 201 Form", 0)
	line = Side(style="thin", color="CED4D9")
	def safe(value):
		return "'" + value if isinstance(value, str) and value.startswith(("=", "+", "-", "@")) else value
	def band(text):
		index = 1 if sheet.max_row == 1 and sheet.cell(1,1).value is None else sheet.max_row + 1
		sheet.merge_cells(start_row=index, start_column=1, end_row=index, end_column=5)
		cell = sheet.cell(index, 1, text)
		cell.font = Font(bold=True, color="30373D", size=11)
		cell.fill = PatternFill("solid", fgColor="EFF1F3")
		sheet.row_dimensions[index].height = 25
	def meta(label, value):
		index = sheet.max_row + 1
		sheet.merge_cells(start_row=index, start_column=1, end_row=index, end_column=2)
		sheet.cell(index, 1, label)
		sheet.merge_cells(start_row=index, start_column=3, end_row=index, end_column=5)
		sheet.cell(index, 3, safe(str(value)) if value not in (None, "") else "Not configured")
		sheet.row_dimensions[index].height = 28 if label in ("Taxable Person Address", "Legal Name (English)") else 23
	band("UAE VAT Return - VAT 201")
	band(form["status"])
	band("Taxable Person Details")
	meta("Company", h.get("company"))
	for label, key in (("TRN", "trn"), ("Legal Name (English)", "legal_name_english"),
		("Legal Name (Arabic)", "legal_name_arabic"), ("Taxable Person Address", "address")):
		meta(label, h.get(key))
	band("VAT Return Period")
	meta("Tax Period From / To", str(h.get("from_date", "")) + " to " + str(h.get("to_date", "")))
	meta("VAT Return Due Date", h.get("due_date"))
	meta("Tax Year End", h.get("tax_year_end"))
	for title, rows, input_section in (("VAT on Sales and All Other Outputs", form["outputs"], False),
		("VAT on Expenses and All Other Inputs", form["inputs"], True)):
		if input_section:
			sheet.row_breaks.append(Break(id=sheet.max_row))
		band(title)
		sheet.append(["Box", "Description", "Amount (AED)", "Recoverable VAT Amount (AED)" if input_section else "VAT Amount (AED)", "Adjustments (AED)"])
		sheet.row_dimensions[sheet.max_row].height = 32
		for row in rows:
			sheet.append([row["box"], row["description"]] + [cell["value"] if cell["applicable"] else None for cell in row["cells"]])
			index = sheet.max_row
			sheet.row_dimensions[index].height = 35 if len(row["description"]) > 65 else 28
			for col, cell in enumerate(row["cells"], 3):
				if cell["provisional"] and cell["applicable"]:
					sheet.cell(index, col).comment = Comment("Known subtotal; return allocation is incomplete. See Detailed Audit export.", "VAT 201")
			if row["total"]:
				for cell in sheet[index]:
					cell.font = Font(bold=True)
					cell.fill = PatternFill("solid", fgColor="EFF1F3")
	band("Net VAT Due")
	for row in form["net"]:
		index = sheet.max_row + 1
		sheet.merge_cells(start_row=index, start_column=1, end_row=index, end_column=4)
		sheet.cell(index, 1, row["box"] + "  " + row["description"])
		sheet.cell(index, 5, row["cells"][0]["value"])
		sheet.row_dimensions[index].height = 27
		if row["cells"][0]["provisional"]:
			sheet.cell(index, 5).comment = Comment("Provisional known subtotal; not a final filing liability.", "VAT 201")
	band("Additional Reporting Requirements")
	meta("Profit Margin Scheme", h.get("profit_margin_scheme"))
	band("Declaration and Authorised Signatory")
	for label, key in (("Name (English)", "signatory_english"), ("Name (Arabic)", "signatory_arabic"),
		("Phone / Mobile", "mobile"), ("Email", "email")):
		meta(label, h.get(key))
	meta("Preparation status", "ERP-generated preparation; not submitted to the FTA.")
	for row in sheet:
		for cell in row:
			cell.border = Border(bottom=line)
			cell.alignment = Alignment(horizontal="right" if isinstance(cell.value, (int, float)) else "left", vertical="center", wrap_text=True)
			if not cell.font.bold:
				cell.font = Font(name="Calibri", size=10, color="30373D")
			if isinstance(cell.value, (int, float)):
				cell.number_format = '#,##0.00;[Red](#,##0.00)'
	for key, width in dict(A=7, B=60, C=20, D=23, E=20).items():
		sheet.column_dimensions[key].width = width
	sheet.sheet_view.showGridLines = False
	sheet.freeze_panes = "C14"
	sheet.print_area = sheet.dimensions
	sheet.print_title_rows = "1:2"
	sheet.sheet_properties.pageSetUpPr.fitToPage = True
	sheet.page_setup.orientation = "portrait"
	sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
	sheet.page_setup.fitToWidth = 1
	sheet.page_setup.fitToHeight = 0
	sheet.oddFooter.center.text = "ERP preparation only - not submitted | Page &P of &N"
	return sheet


def workbook_bytes(result, audit=False, audit_bytes=None):
	import io
	from openpyxl import Workbook, load_workbook
	workbook = load_workbook(io.BytesIO(audit_bytes)) if audit else Workbook()
	if not audit:
		workbook.remove(workbook.active)
	form_sheet(workbook, result)
	workbook.active = 0
	output = io.BytesIO()
	workbook.save(output)
	return output.getvalue()
