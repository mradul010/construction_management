import io
import unittest
from datetime import date
from unittest.mock import patch

import frappe
from construction_management.construction_management.report.uae_vat_201_return import engine
from construction_management.construction_management.report.uae_vat_201_return.uae_vat_201_return import workbook_bytes, export_rows


class TestUAEVAT201Return(unittest.TestCase):
	def setUp(self):
		self.enterContext(patch.object(engine.frappe,"get_precision",return_value=2))

	def source(self, **values):
		return dict(dict(voucher_type="Sales Invoice",voucher_no="TEST-SI",source_row="TEST-TAX",posting_date="2026-06-01",base=1000,raw_vat=50,tax_rate=5,tax_account="TEST-VAT",company="TEST",currency="AED"),**values)

	def rule(self, box="1b", **values):
		return frappe._dict(dict(dict(box=box,approved=1,approved_by="test-reviewer",tax_date="2026-06-01",reason="Isolated fixture - reviewed place of supply",supporting_reference="TEST ONLY",recovery_confirmed=1,recovery_percent=100,is_adjustment=0,override_base=0),**values))

	def event(self, source=None, rule=None):
		event,error=engine.apply_rule(source or self.source(),rule or self.rule(),"AED")
		self.assertIsNone(error)
		return event

	def test_all_boxes_independently_and_totals(self):
		events=[]
		for box in engine.OUTPUT+engine.INPUT:
			vat=0 if box in ("4","5") else -5 if box=="2" else 50
			source=self.source(raw_vat=vat,voucher_type="Purchase Invoice" if box in ("3","9","10") else "Sales Invoice",reverse_charge="Y")
			rule=self.rule(box,override_base=1,taxable_base_aed=-100 if box=="2" else 1000)
			event=self.event(source,rule)
			boxes={b["box"]:b for b in engine.aggregate([event]).values()}
			self.assertEqual(boxes[box]["vat_amount"],vat)
			self.assertEqual(sum(b["vat_amount"] for k,b in boxes.items() if k in engine.OUTPUT+engine.INPUT),vat)
			events.append(event)
		boxes=engine.aggregate(events)
		self.assertEqual(boxes["12"]["amount"],495)
		self.assertEqual(boxes["13"]["amount"],100)
		self.assertEqual(boxes["14"]["amount"],395)

	def test_seven_emirates(self):
		for box,emirate in zip(engine.OUTPUT[:7],engine.get_emirates()):
			self.assertEqual(self.event(rule=self.rule(box))["emirate"],emirate)

	def test_sales_return_and_purchase_return_keep_signs(self):
		for dt,box in (("Sales Invoice","1b"),("Purchase Invoice","9")):
			e=self.event(self.source(voucher_type=dt,base=-1000,raw_vat=-50),self.rule(box))
			self.assertEqual((e["amount"],e["vat_amount"]),(-1000,-50))

	def test_zero_rated_and_exempt_are_separate(self):
		for box in ("4","5"):
			e=self.event(self.source(raw_vat=0),self.rule(box))
			self.assertEqual((e["box"],e["amount"],e["vat_amount"]),(box,1000,0))
			self.assertIsNotNone(engine.apply_rule(self.source(),self.rule(box),"AED")[1])

	def test_reverse_charge_due_and_partial_recovery(self):
		source=self.source(voucher_type="Purchase Invoice",reverse_charge="Y")
		events=[self.event(source,self.rule("3")),self.event(source,self.rule("10",recovery_percent=40))]
		boxes=engine.aggregate(events)
		self.assertEqual((boxes["12"]["amount"],boxes["13"]["amount"],boxes["14"]["amount"]),(50,20,30))
		self.assertIsNotNone(engine.apply_rule(dict(source,reverse_charge="N"),self.rule("10"),"AED")[1])

	def test_imports_require_customs_base(self):
		for box in ("6","7"):
			self.assertIsNotNone(engine.apply_rule(self.source(),self.rule(box),"AED")[1])
			e=self.event(rule=self.rule(box,override_base=1,taxable_base_aed=2000))
			self.assertEqual(e["amount"],2000)

	def test_nonrecoverable_and_partial_input_vat(self):
		for percent,expected in ((0,0),(30,15),(100,50)):
			e=self.event(self.source(voucher_type="Purchase Invoice"),self.rule("9",recovery_percent=percent))
			self.assertEqual(e["vat_amount"],expected)
			self.assertEqual(e["nonrecoverable"],50-expected)
			self.assertEqual(e["amount"],1000 if percent else 0)
		self.assertIsNotNone(engine.apply_rule(self.source(voucher_type="Purchase Invoice"),self.rule("9",recovery_confirmed=0),"AED")[1])

	def test_foreign_invoice_uses_base_amount_not_document_currency(self):
		e=self.event(self.source(currency="USD",base=3672.5,raw_vat=183.63))
		self.assertEqual((e["amount"],e["vat_amount"]),(3672.5,183.63))

	def test_non_aed_company_requires_approved_conversion(self):
		self.assertIsNotNone(engine.apply_rule(self.source(),self.rule(),"USD")[1])
		e,error=engine.apply_rule(self.source(),self.rule(aed_exchange_rate=3.6726),"USD")
		self.assertIsNone(error)
		self.assertEqual((e["amount"],e["vat_amount"]),(3672.6,183.63))

	def test_adjustment_uses_posted_tax_and_its_own_column(self):
		e=self.event(self.source(voucher_type="Journal Entry",base=None,debit=25,credit=0),self.rule(is_adjustment=1,adjustment_type="Bad Debt Relief",override_base=1,taxable_base_aed=0))
		self.assertEqual((e["amount"],e["vat_amount"],e["adjustment"]),(0,0,-25))
		self.assertEqual(engine.aggregate([e])["12"]["amount"],-25)

	def test_ordinary_correction_is_not_an_adjustment_column_entry(self):
		self.assertIsNotNone(engine.apply_rule(self.source(),self.rule(is_adjustment=1,adjustment_type="VAT Correction"),"AED")[1])
		self.assertEqual(self.event(rule=self.rule(adjustment_type="VAT Correction"))["vat_amount"],50)

	def test_negative_net_is_not_clamped(self):
		e=self.event(self.source(voucher_type="Purchase Invoice"),self.rule("9"))
		self.assertEqual(engine.aggregate([e])["14"]["amount"],-50)

	def test_missing_approval_evidence_is_not_a_tax_classification(self):
		for field in ("approved_by","reason","supporting_reference"):
			rule=self.rule(**{field:None})
			self.assertIsNotNone(engine.apply_rule(self.source(),rule,"AED")[1])

	def test_monthly_leap_year_and_non_calendar_quarter(self):
		settings=frappe._dict(custom_vat201_periods=[frappe._dict(frequency="Quarterly",quarter_number=1,from_date="2026-02-01",to_date="2026-04-30"),frappe._dict(frequency="Quarterly",quarter_number=2,from_date="2026-05-01",to_date="2026-07-31")])
		f=engine.normalize(dict(year=2024,month=2,period_type="Monthly"),settings)
		self.assertEqual((f.from_date,f.to_date),(date(2024,2,1),date(2024,2,29)))
		f=engine.normalize(dict(year=2026,quarter=2,period_type="Quarterly"),settings)
		self.assertEqual((f.from_date,f.to_date),(date(2026,5,1),date(2026,7,31)))
		with self.assertRaises(frappe.ValidationError):
			engine.normalize(dict(year=2026,quarter=1,period_type="Quarterly"),frappe._dict())

	def test_custom_date_range_and_empty_totals(self):
		f=engine.normalize(dict(from_date="2026-05-10",to_date="2026-06-09"),frappe._dict())
		self.assertEqual((f.from_date,f.to_date),(date(2026,5,10),date(2026,6,9)))
		self.assertTrue(all(b["amount"]==0 and b["vat_amount"]==0 for b in engine.aggregate([]).values()))

	def test_three_decimal_precision(self):
		with patch.object(engine.frappe,"get_precision",return_value=3):
			self.assertEqual(self.event(self.source(raw_vat="1.2344"))["vat_amount"],1.234)

	def test_actual_vat_is_never_recomputed_as_five_percent(self):
		self.assertEqual(self.event(self.source(raw_vat="39.14",tax_rate=0))["vat_amount"],39.14)

	def test_approved_non_vat_expenses_are_auditable_exclusions(self):
		e=self.event(rule=self.rule("Outside VAT"))
		self.assertEqual(e["classification"],"Outside VAT scope")
		self.assertEqual(engine.aggregate([e])["14"]["amount"],0)
		self.assertEqual(e["supporting_reference"],"TEST ONLY")

	def test_reverse_charge_ledger_sides_and_return(self):
		row=frappe._dict(name="GL-TEST",account="VAT",voucher_type="Purchase Invoice",voucher_no="PI-TEST",debit=50,credit=50)
		headers={("Purchase Invoice","PI-TEST"):dict(reverse_charge="Y")}
		self.assertEqual(engine.ledger_sides(row,{"VAT":"Both"},headers,[]),[("Input",50),("Output",50)])
		headers[("Purchase Invoice","PI-TEST")]["is_return"]=1
		self.assertEqual(engine.ledger_sides(row,{"VAT":"Both"},headers,[]),[("Output",-50),("Input",-50)])

	def test_shared_account_is_merged_without_duplicate_or_last_row_side(self):
		rows = [frappe._dict(account="VAT", custom_vat201_side=side) for side in ("Output", "Input", "Input")]
		self.assertEqual(engine.account_mappings(rows), {"VAT": "Both"})
		for dt, debit, credit, expected in (
			("Sales Invoice", 0, 50, ("Output", 50)),
			("Sales Invoice", 50, 0, ("Output", -50)),
			("Purchase Invoice", 50, 0, ("Input", 50)),
			("Purchase Invoice", 0, 50, ("Input", -50)),
			("Journal Entry", 0, 50, ("Unallocated", -50)),
		):
			row = frappe._dict(account="VAT", voucher_type=dt, debit=debit, credit=credit)
			self.assertEqual(engine.ledger_sides(row, engine.account_mappings(rows), {}, []), [expected])

	def test_first_return_default_and_cross_year_cycle(self):
		settings = frappe._dict(custom_vat201_periods=[
			frappe._dict(from_date="2026-05-01", to_date="2026-07-31", frequency="Quarterly", quarter_number=2),
			frappe._dict(from_date="2026-11-01", to_date="2027-01-31", frequency="Quarterly", quarter_number=4),
		])
		f = engine.normalize({}, settings)
		self.assertEqual((f.from_date, f.to_date), (date(2026,5,1), date(2026,7,31)))
		f = engine.normalize(dict(period_type="Quarterly", year=2026, quarter=4), settings)
		self.assertEqual((f.from_date, f.to_date), (date(2026,11,1), date(2027,1,31)))
		f = engine.normalize(dict(from_date="2026-08-01", to_date="2026-10-31"), settings)
		self.assertEqual((f.from_date, f.to_date), (date(2026,8,1), date(2026,10,31)))

	def test_taxpayer_uses_registered_legal_name_not_company_rename(self):
		settings = frappe._dict(custom_vat201_legal_name_english="Registered Name OPC",
			custom_vat201_registration_date="2026-05-01", custom_vat201_registered_emirate="Ras Al Khaimah",
			custom_vat201_periods=[frappe._dict(from_date="2026-05-01", to_date="2026-07-31", frequency="Quarterly", quarter_number=2, profit_margin_scheme="No", due_date="2026-08-31")])
		with patch.object(frappe,"get_doc", return_value=frappe._dict(company_name="ERP Name", tax_id="105452509000003", default_currency="AED")), patch.object(frappe,"get_all", return_value=[]):
			header = engine.taxpayer("ERP Name", settings, engine.normalize({},settings))
		self.assertEqual(header["legal_name_english"], "Registered Name OPC")
		self.assertEqual(header["registered_emirate"], "Ras Al Khaimah")
		self.assertEqual(header["trn"], "105452509000003")
		self.assertEqual(header["filing_frequency"], "Quarterly")
		self.assertEqual(header["profit_margin_scheme"], "No")
		self.assertEqual(header["due_date"], "2026-08-31")

	def test_profit_margin_scheme_is_blocked(self):
		with patch.object(engine,"access"),patch.object(engine,"settings_for",return_value=frappe._dict()),patch.object(engine,"taxpayer",return_value={"profit_margin_scheme":"Yes"}):
			with self.assertRaises(frappe.ValidationError):
				engine.build(dict(company="TEST",from_date="2026-01-01",to_date="2026-01-31"))

	def test_access_is_checked_before_financial_queries(self):
		with patch.object(engine,"access",side_effect=frappe.PermissionError),patch.object(engine,"source_data") as source:
			with self.assertRaises(frappe.PermissionError):
				engine.build(dict(company="TEST"))
			source.assert_not_called()

	def test_excel_contains_metadata_details_reviews_and_reconciliation(self):
		from openpyxl import load_workbook
		result=dict(header={"company":"Test Company","trn":"TEST","from_date":"2026-06-01"},status="Needs Review",warnings=["=UNTRUSTED()"],boxes=list(engine.aggregate([self.event()]).values()),events=[self.event()],reviews=[],reconciliation=[dict(description="Output VAT",report_amount=50,ledger_amount=49,difference=1)])
		book=load_workbook(io.BytesIO(workbook_bytes(result)))
		values=[c.value for sheet in book for row in sheet for c in row]
		for text in ("Test Company","VAT 201 Preparation","Transactions","Reconciliation","Review"):
			if text in book.sheetnames:
				continue
			self.assertIn(text,values)
		self.assertIn("'=UNTRUSTED()",values)
		self.assertFalse(any(c.data_type=="f" for row in book.active for c in row))
