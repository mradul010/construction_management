import csv
import io
from decimal import Decimal

import frappe
from frappe import _
from frappe.utils import flt

from construction_management.construction_management.report.uae_vat_201_return.automatic import build, detail_remarks
from construction_management.construction_management.report.uae_vat_201_return.engine import OUTPUT, INPUT


def get_columns():
	return [dict(fieldname="box",label=_("Box Number"),fieldtype="Data",width=100),
		dict(fieldname="description",label=_("Description"),fieldtype="Data",width=440)] + [
		dict(fieldname=key,label=_(label),fieldtype="Currency",options="currency",width=170,align="right")
		for key,label in (("amount","Amount (AED)"),("vat_amount","VAT Amount (AED)"),("adjustment","Adjustments (AED)"))
	] + [dict(fieldname="currency",fieldtype="Data",hidden=1)]


def execute(filters=None):
	result = build(filters)
	frappe.local.vat201_current_result = result
	rows = [dict(box,_box=True) for box in result["boxes"]]
	from construction_management.construction_management.report.uae_vat_201_return.form_layout import form_html
	message = form_html(result, interactive=True)
	return get_columns(),rows,message,None,summary_cards(result),True


def summary_cards(result):
	boxes = {b["box"]:b for b in result["boxes"]}
	def pending(box, field="amount"):
		return boxes[box].get("completeness", {}).get(field, {}).get("provisional", False)
	def card(label, value, provisional=False):
		return dict(label=_(label),
			value=_("Not Available") if value is None else value,
			datatype="Data" if value is None else "Currency",currency="AED",precision=2,provisional=provisional)
	return [card("Total Output Tax - Box 12",boxes["12"]["amount"],pending("12")),
		card("Total Recoverable Input Tax - Box 13",boxes["13"]["amount"],pending("13")),
		card("Net VAT Refundable - Box 14" if boxes["14"]["amount"] < 0 else "Net VAT Payable / Refundable - Box 14",boxes["14"]["amount"],pending("14"))]


@frappe.whitelist()
def period_context(filters):
	from construction_management.construction_management.report.uae_vat_201_return.engine import access, settings_for, normalize, taxpayer
	f = frappe._dict(frappe.parse_json(filters) or {})
	access(f.company)
	settings = settings_for(f.company)
	f = normalize(f, settings)
	header = taxpayer(f.company,settings,f)
	header["available_periods"] = [f"{p.from_date} to {p.to_date}" for p in settings.get("custom_vat201_periods", [])]
	return header


@frappe.whitelist()
def report_context(filters):
	result = build(filters)
	return context_for(result)


def context_for(result):
	# Source permissions are checked by the reader; no stored evidence snapshots are returned.
	context = {key:result.get(key) for key in ("header","status","calculation_description","return_complete","warnings","details","reconciliation","observed","unallocated_output","review_summary","audit_rows","tax_measures","ledger_summary","recorded_breakdown","measure_labels","settings_name","calculated_at")}
	from construction_management.construction_management.report.uae_vat_201_return.form_layout import form_html
	context["form_print_html"] = form_html(result)
	return context


@frappe.whitelist()
def transactions(filters, box, fieldname=None):
	result = build(filters)
	keys = OUTPUT if box in ("8","12") else INPUT if box in ("11","13") else OUTPUT+INPUT if box=="14" else [box]
	events = [dict(e) for e in result["events"] if e["box"] in keys]
	for event in events:
		if box in ("12","13","14") and fieldname in (None,"amount"):
			value = flt(event["vat_amount"])+flt(event["adjustment"])
			if box=="14" and event["box"] in INPUT:
				value = -value
		elif box in ("12","13","14"):
			value = 0
		else:
			value = flt(event.get(fieldname or "vat_amount"))
		event["contribution"] = value
		event["display_remarks"] = detail_remarks(event)
	return events


def export_rows(result):
	rows = [["UAE VAT 201 RETURN",result["status"]],["ERP-generated preparation report; not submitted to the FTA."]]
	rows += [[key.replace("_"," ").title(),value if value not in (None,"") else "Not Configured"] for key,value in result["header"].items()]
	rows += [["Warnings"]]+[[w] for w in result["warnings"]]
	rows += [["VAT 201 Summary"],["Box","Description","Amount (AED)","VAT Amount (AED)","Adjustment (AED)"]]
	rows += [[b["box"],b["description"]]+[b[k] for k in ("amount","vat_amount","adjustment")] for b in result["boxes"]]
	rows += [["Box Component Completeness"],["Box","Component","Status","Reason"]]
	for box in result["boxes"]:
		for component, state in box.get("completeness", {}).items():
			rows.append([box["box"],component,"Provisional known subtotal" if state["provisional"] else "Complete", "; ".join(state["reasons"])])
	if result.get("observed"):
		rows += [["Report Totals"],["Description","Amount (AED)","Status"]]
		rows += [[card["label"],card["value"],"Provisional known subtotal" if card["provisional"] else "Complete" if card["datatype"]=="Currency" else "Unavailable"] for card in summary_cards(result)]
	fields = ("posting_date","voucher_no","party","invoice_taxable_amount","recorded_vat_amount","treatment","vat_box","display_remarks")
	rows += [["Transaction Details"],["Date","Voucher Number","Party","Taxable Amount","VAT Amount","VAT Treatment","VAT Box","Remarks"]]
	for row in result.get("details",result["events"]):
		values = dict(row)
		values.setdefault("invoice_taxable_amount",row.get("base"))
		values.setdefault("recorded_vat_amount",row.get("raw_vat"))
		values.setdefault("vat_box",row.get("box"))
		values.setdefault("display_remarks",row.get("remarks"))
		rows.append([values.get(k) for k in fields])
	rows += [["Source Notes"],["Voucher Type","Voucher","Source Row","Source limitation","Recorded VAT (VAT-role signs)","Source context","Raw GL VAT (debit - credit)"]]
	rows += [[r.get("voucher_type"),r.get("voucher_no"),r.get("source_row"),r.get("issue"),r.get("recorded_vat_amount",r.get("raw_vat")),r.get("recommended_action"),
		flt(r.get("debit"))-flt(r.get("credit")) if "debit" in r else None] for r in result["reviews"]]
	rows += [["Tax Reconciliation"],["Description","Included VAT","Ledger VAT","Difference","Explanation"]]
	rows += [[r.get(k) for k in ("description","report_amount","ledger_amount","difference","explanation")] for r in result["reconciliation"]]
	audit_fields = ("voucher_type","voucher_no","vat_account","posting_dates","tax_dates","source_rows","base_rows","gl_rows","raw_currency","raw_debit","raw_credit","raw_movement","recorded_output","recorded_input","claimed_output","claimed_input","known_nonclaimed","pending_output","pending_input","timing_output","timing_input","residual_output","residual_input","classification_difference","unallocated_gl","control_movement","source_to_gl_difference","category","issue")
	rows += [["Voucher Account Reconciliation"],[key.replace("_"," ").title() for key in audit_fields]]
	rows += [[row.get(key) for key in audit_fields] for row in result.get("audit_rows", [])]
	rows += [["Ledger Balances"],["Measure","Amount","Currency"]]
	rows += [[key.replace("_"," ").title(),value,result["ledger_summary"]["currency"]] for key,value in result.get("ledger_summary",{}).items() if key!="currency"]
	rows += [["Recorded Tax Bridge"],["Measure","Amount (AED)"]]
	rows += [[result.get("measure_labels",{}).get(key,key.replace("_"," ").title()),value] for key,value in result.get("tax_measures",{}).items()]
	breakdown_fields = ("voucher_type","category","raw_debit","raw_credit","recorded_output","recorded_input","claimed_output","claimed_input","pending_output","pending_input","known_nonclaimed","residual_output","residual_input")
	rows += [["Recorded VAT by Voucher Type"],[key.replace("_"," ").title() for key in breakdown_fields]]
	rows += [[row.get(key) for key in breakdown_fields] for row in result.get("recorded_breakdown",[])]
	decision_fields = ("voucher_type","voucher_no","source_row","tax_account","posting_date","tax_date","recovery_source_amount","recovery_percent","recoverable_vat_amount","recovery_status","treatment_source","policy_id","policy_name","policy_approved_by","policy_approval_time","approved_by","approval_time","reason","supporting_reference")
	rows += [["Recovery Decisions"],[key.replace("_"," ").title() for key in decision_fields]]
	rows += [[row.get(key) for key in decision_fields] for row in result.get("details",[]) if row.get("recovery_status")]
	rows += [["Recorded source amounts"]]+[[k,v] for k,v in result.get("observed",{}).items()]
	rows += [["VAT Control Movement (debit - credit)",result.get("control_movement")]]
	rows += [["Invoice tax versus GL"],["Voucher Type","Voucher","VAT Account","Difference","Currency"]]
	rows += [[r.get(k) for k in ("voucher_type","voucher_no","tax_account","difference","currency")] for r in result.get("invoice_checks",[])]
	return [[("'"+v if v.startswith(("=","+","-","@")) else v) if isinstance(v,str) else v for v in row] for row in rows]


def audit_workbook_bytes(result):
	from openpyxl import Workbook
	from openpyxl.styles import Alignment, Font, PatternFill
	from openpyxl.comments import Comment
	from frappe.utils import getdate
	workbook = Workbook()
	precision = frappe.get_precision("GL Entry","debit",currency="AED")
	pattern = '#,##0' + ('.'+'0'*precision if precision else '')
	sections = {name:[] for name in ("VAT 201 Summary","Box Component Completeness","Report Totals","Transaction Details","Source Notes","Tax Reconciliation","Voucher Account Reconciliation","Recorded source amounts","Invoice tax versus GL","Ledger Balances","Recorded Tax Bridge","Recorded VAT by Voucher Type","Recovery Decisions")}
	current = None
	metadata = []
	for row in export_rows(result):
		if len(row)==1 and row[0] in sections:
			current = row[0]
		elif current:
			sections[current].append(row)
		else:
			metadata.append(row)
	main = workbook.active
	main.title = "VAT 201 Preparation"
	for row in [["UAE VAT 201 RETURN", result["status"]], ["Company", result["header"].get("company")],
		["TRN",result["header"].get("trn")], ["Period",result["header"].get("from_date"),result["header"].get("to_date")],
		["Preparation only; not submitted to the FTA."], [], *sections["VAT 201 Summary"], [], *sections["Report Totals"]]:
		main.append(row)
	main.freeze_panes = "C8"
	for index, box in enumerate(result["boxes"], start=8):
		for col, field in enumerate(("amount","vat_amount","adjustment"),start=3):
			state = box.get("completeness",{}).get(field,{})
			if state.get("provisional"):
				main.cell(index,col).fill = PatternFill("solid",fgColor="FFF3CD")
				main.cell(index,col).comment = Comment("Provisional known subtotal. " + "; ".join(state["reasons"]),"VAT 201")
	for title, rows in (("Transactions",sections["Transaction Details"]),("Source Notes",sections["Source Notes"]),
		("Completeness",sections["Box Component Completeness"]),("Reconciliation",sections["Tax Reconciliation"]),
		("Source Ledger Audit",sections["Voucher Account Reconciliation"]),("Invoice GL Checks",sections["Invoice tax versus GL"]),
		("Taxpayer and Notes",metadata + sections["Recorded source amounts"]),
		("Ledger Balances",sections["Ledger Balances"]),("Recorded Tax Bridge",sections["Recorded Tax Bridge"]),
		("VAT by Voucher Type",sections["Recorded VAT by Voucher Type"]),("Recovery Decisions",sections["Recovery Decisions"])):
		sheet = workbook.create_sheet(title)
		for row in rows:
			sheet.append(row)
		sheet.freeze_panes = "D2" if title in ("Transactions","Review","Source Ledger Audit") else "B2"
		if sheet.max_row > 1 and title != "Taxpayer and Notes":
			sheet.auto_filter.ref = sheet.dimensions
	transactions = workbook["Transactions"]
	for index in range(2,transactions.max_row+1):
		cell = transactions.cell(index,1)
		if cell.value:
			cell.value = getdate(cell.value)
			cell.number_format = "dd-mm-yyyy"
	for index, source in enumerate(result.get("details",result["events"]),start=2):
		if source.get("voucher_type") and source.get("voucher_no"):
			cell = transactions.cell(index,2)
			cell.hyperlink = frappe.utils.get_url_to_form(source["voucher_type"],source["voucher_no"])
			cell.style = "Hyperlink"
	for sheet in workbook:
		for cell in sheet[1]:
			cell.font = Font(bold=True,color="213547")
			cell.fill = PatternFill("solid",fgColor="EDF2F5")
		for column in sheet.columns:
			sheet.column_dimensions[column[0].column_letter].width = min(60,max(16,max(len(str(c.value or "")) for c in column)+2))
			for cell in column:
				if isinstance(cell.value,str) and cell.value.startswith(("=","+","-","@")):
					cell.value = "'" + cell.value
				cell.alignment = Alignment(horizontal="right" if isinstance(cell.value,(int,float,Decimal)) else "left",vertical="top",wrap_text=not isinstance(cell.value,(int,float,Decimal)))
				if isinstance(cell.value,(int,float,Decimal)):
					cell.number_format = pattern+';[Red]('+pattern+')'
		if sheet == main:
			for cell in sheet[7]:
				cell.font = Font(bold=True)
	for key, width in {"A":16,"B":28,"C":34,"D":20,"E":20,"F":34,"G":14,"H":60}.items():
		transactions.column_dimensions[key].width = width
	output = io.BytesIO()
	workbook.save(output)
	return output.getvalue()


def workbook_bytes(result, detailed=False):
	from construction_management.construction_management.report.uae_vat_201_return.form_layout import workbook_bytes as render
	return render(result, audit=detailed, audit_bytes=audit_workbook_bytes(result) if detailed else None)


def pdf_bytes(result, detailed=False):
	from frappe.utils.pdf import get_pdf
	from construction_management.construction_management.report.uae_vat_201_return.form_layout import form_html
	content = frappe.render_template("construction_management/report/uae_vat_201_return/return_print.html",dict(result=result,precision=frappe.get_precision("GL Entry","debit",currency="AED"))) if detailed else "<html><head><meta charset='utf-8'></head><body>" + form_html(result) + "</body></html>"
	return get_pdf(content,options={"page-size":"A4","margin-top":"10mm","margin-bottom":"10mm","margin-left":"10mm","margin-right":"10mm","print-media-type":None})


@frappe.whitelist()
def download(filters, file_type="xlsx", detailed=0):
	result = build(filters)
	if file_type not in ("xlsx","csv","pdf"):
		frappe.throw(_("Unsupported export format"))
	frappe.has_permission("GL Entry","print" if file_type=="pdf" else "export",throw=True)
	detailed = bool(frappe.utils.cint(detailed))
	if file_type=="xlsx":
		content = workbook_bytes(result, detailed=detailed)
	elif file_type=="pdf":
		content = pdf_bytes(result, detailed=detailed)
	else:
		stream = io.StringIO()
		csv.writer(stream).writerows(export_rows(result))
		content = stream.getvalue().encode("utf-8-sig")
	frappe.local.response.filename = f"UAE-VAT-201-{result['header']['to_date']}{'-Detailed-Audit' if detailed else ''}.{file_type}"
	frappe.local.response.filecontent = content
	frappe.local.response.type = "pdf" if file_type=="pdf" else "binary"
