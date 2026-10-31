import copy
import io
import unittest
from unittest.mock import patch

from lxml import html
from openpyxl import load_workbook

from construction_management.construction_management.report.uae_vat_201_return import engine, form_layout
from construction_management.construction_management.report.uae_vat_201_return import uae_vat_201_return as report


class TestVAT201FormLayout(unittest.TestCase):
	def setUp(self):
		self.enterContext(patch.object(engine.frappe,"get_precision",return_value=2))
		self.result=dict(header=dict(company="Test Company",legal_name_english="Test Company Legal",trn="100000000000001",
			from_date="2026-05-01",to_date="2026-07-31"),status="Calculated from saved ERPNext records - incomplete return allocation",
			boxes=list(engine.aggregate([dict(box="1f",amount=1000,vat_amount=50,adjustment=0),dict(box="9",amount=200,vat_amount=10,adjustment=0)]).values()),
			events=[],details=[],reviews=[],warnings=[],reconciliation=[])

	def test_exact_three_cards_use_return_boxes_not_recorded_ledger(self):
		self.result["tax_measures"]=dict(recorded_output=999,recorded_input=888)
		cards=report.summary_cards(self.result)
		self.assertEqual(len(cards),3)
		self.assertEqual([card["value"] for card in cards],[50,10,40])
		self.assertTrue(all(card["precision"]==2 and card["currency"]=="AED" for card in cards))
		next(box for box in self.result["boxes"] if box["box"]=="14")["amount"]=-40
		self.assertEqual(report.summary_cards(self.result)[2]["label"],"Net VAT Refundable - Box 14")

	def test_all_statutory_rows_and_native_column_distinctions(self):
		before=copy.deepcopy(self.result)
		form=form_layout.form_data(self.result)
		self.assertEqual([r["box"] for r in form["outputs"]],form_layout.OUTPUT_KEYS)
		self.assertEqual([r["box"] for r in form["inputs"]],["9","10","11"])
		self.assertEqual([r["box"] for r in form["net"]],["12","13","14"])
		root=html.fromstring(form_layout.form_html(self.result,interactive=True))
		self.assertEqual(len(root.xpath('//*[@data-vat-box]')),20)
		self.assertIn("Recoverable VAT Amount (AED)",root.text_content())
		for box in ("4","5"):
			self.assertEqual(root.xpath('//*[@data-vat-box="'+box+'"]//*[@data-value-field="vat_amount"]')[0].text_content().strip(),"")
		self.assertTrue(root.xpath('//*[@data-vat-box="12"]//*[@data-field="amount"]'))
		self.assertFalse(root.xpath('//*[@data-vat-box="12"]//*[@data-value-field="vat_amount"]'))
		self.assertEqual(self.result,before)

	def test_unexpected_nonzero_is_never_discarded_as_not_applicable(self):
		box=next(b for b in self.result["boxes"] if b["box"]=="4")
		box["vat_amount"]=3.25
		form=form_layout.form_data(self.result)
		row=next(r for r in form["outputs"] if r["box"]=="4")
		self.assertTrue(row["cells"][1]["applicable"])
		self.assertEqual(row["cells"][1]["value"],3.25)

	def test_default_excel_is_form_only_numeric_and_print_ready(self):
		book=load_workbook(io.BytesIO(report.workbook_bytes(self.result)))
		self.assertEqual(book.sheetnames,["VAT 201 Form"])
		sheet=book.active
		rows=list(sheet.values)
		self.assertEqual(next(r for r in rows if r[0]=="1f")[2:5],(1000,50,0))
		self.assertEqual(next(r for r in rows if r[0]=="4")[2:5],(0,None,None))
		self.assertEqual(next(r for r in rows if isinstance(r[0],str) and r[0].startswith("14  "))[4],40)
		self.assertTrue(sheet.print_area)
		self.assertEqual(sheet.page_setup.fitToWidth,1)
		self.assertEqual(sheet.page_setup.orientation,"portrait")
		self.assertTrue(all(c.number_format.startswith("#,##0.00") for row in sheet for c in row if isinstance(c.value,(int,float))))

	def test_audit_remains_explicit_and_preserves_all_supporting_sheets(self):
		book=load_workbook(io.BytesIO(report.workbook_bytes(self.result,detailed=True)))
		self.assertEqual(book.sheetnames[0],"VAT 201 Form")
		for name in ("Transactions","Source Ledger Audit","Recovery Decisions","Completeness","Recorded Tax Bridge"):
			self.assertIn(name,book.sheetnames)

	def test_html_escapes_source_values_and_preserves_provisional_marker(self):
		self.result["header"]["legal_name_english"]="<script>source</script>"
		self.result["boxes"][0]["completeness"]={"vat_amount":dict(provisional=True)}
		content=form_layout.form_html(self.result)
		self.assertNotIn("<script>source</script>",content)
		self.assertIn("&lt;script&gt;source&lt;/script&gt;",content)
		self.assertIn("not a final filing liability",content)
