import unittest
from contextlib import ExitStack
from datetime import date
from io import BytesIO
from unittest.mock import patch

import frappe
from construction_management.construction_management.report.project_profitability import project_profitability as report


class TestProjectProfitability(unittest.TestCase):
	def test_expense_classification_inherits_without_guessing(self):
		accounts = [
			frappe._dict(name="COGS", root_type="Expense", account_type="Cost of Goods Sold"),
			frappe._dict(name="Materials", root_type="Expense", parent_account="COGS"),
			frappe._dict(name="Admin", root_type="Expense", custom_profitability_classification="Indirect"),
			frappe._dict(name="Fees", root_type="Expense", parent_account="Admin"),
			frappe._dict(name="Unknown", root_type="Expense"),
			frappe._dict(name="Income", root_type="Income"),
		]
		classes = report.account_classes(accounts)
		self.assertEqual(classes["Materials"], "direct_expenses")
		self.assertEqual(classes["Fees"], "indirect_expenses")
		self.assertEqual(classes["Unknown"], "unclassified_expenses")
		self.assertEqual(classes["Income"], "income")
		accounts[1].custom_profitability_classification = "Indirect"
		self.assertEqual(report.account_classes(accounts)["Materials"], "indirect_expenses")

	def test_april_fiscal_periods_are_contiguous_and_clipped(self):
		def fiscal(day, **kwargs):
			year = day.year if day.month >= 4 else day.year - 1
			return str(year), date(year, 4, 1), date(year + 1, 3, 31)
		with patch.object(report, "get_fiscal_year", side_effect=fiscal):
			for kind, count in (("Monthly", 24), ("Quarterly", 8), ("Yearly", 2)):
				filters = frappe._dict(company="Example", from_date=date(2025, 4, 1), to_date=date(2027, 3, 31), period_type=kind)
				periods = report.get_periods(filters)
				self.assertEqual(len(periods), count)
				self.assertEqual(periods[0].from_date, filters.from_date)
				self.assertEqual(periods[-1].to_date, filters.to_date)
				for previous, current in zip(periods, periods[1:]):
					self.assertEqual((current.from_date - previous.to_date).days, 1)
				if kind == "Quarterly":
					self.assertEqual(periods[0].to_date, date(2025, 6, 30))
				filters.from_date, filters.to_date = date(2025, 5, 10), date(2025, 8, 16)
				clipped = report.get_periods(filters)
				self.assertEqual(clipped[0].from_date, filters.from_date)
				self.assertEqual(clipped[-1].to_date, filters.to_date)

	def test_signs_profits_zero_income_and_currency_precision(self):
		classes = {"I": "income", "D": "direct_expenses", "N": "indirect_expenses", "U": "unclassified_expenses"}
		entries = [frappe._dict(account="I", balance=-100), frappe._dict(account="D", balance=40), frappe._dict(account="N", balance=10)]
		with patch.object(report.frappe, "get_precision", return_value=3):
			values = report.totals(entries, classes, "TEST")
			self.assertEqual((values["income"], values["gross_profit"], values["net_profit"], values["margin"]), (100, 60, 50, 50))
			self.assertEqual(report.totals([frappe._dict(account="D", balance="1.2344")], classes, "TEST")["expenses"], 1.234)
			self.assertIsNone(report.totals(entries[1:], classes, "TEST")["margin"])
			entries.append(frappe._dict(account="U", balance=5))
			values = report.totals(entries, classes, "TEST")
			self.assertIsNone(values["gross_profit"])
			self.assertEqual(values["net_profit"], 45)

	def test_empty_scope_never_reads_unrestricted_gl(self):
		with patch.object(report.frappe.db, "sql") as sql:
			self.assertEqual(report.read_entries(frappe._dict(project="Hidden"), [], ["Account"]), [])
			self.assertEqual(report.read_entries(frappe._dict(), [], []), [])
			sql.assert_not_called()

	def test_drilldown_checks_existing_report_access(self):
		with patch.object(report.frappe, "has_permission"), patch.object(report, "get_report_doc", side_effect=frappe.PermissionError):
			with self.assertRaises(frappe.PermissionError):
				report.drilldown({"company": "Example"}, {"project": "Hidden"})

	def example(self, view="Cost Center Breakdown", show=True, entries=None):
		projects = [frappe._dict(name="P1", project_name="Shared name"), frappe._dict(name="P2", project_name="Shared name")]
		accounts = [frappe._dict(name=name, root_type=root, account_type=kind, lft=index)
			for index, (name, root, kind) in enumerate([
				("I", "Income", ""), ("D1", "Expense", "Cost of Goods Sold"), ("D2", "Expense", "Cost of Goods Sold")])]
		if entries is None:
			entries = [("P1", "Revenue", "I", -1548437.56, 1),
				("P1", "Block", "D1", 10000, 1), ("P1", "Block", "D2", 9497.94, 2),
				("P1", "Other", "D1", 917710.30, 2), ("P2", "Block", "D1", 12, 1)]
		entries = [frappe._dict(project=p, cost_center=c, account=a, balance=b, posting_date=date(2026, m, 15)) for p, c, a, b, m in entries]
		filters = frappe._dict(company="Example", view_type=view, metric="Expenses", show_percentage_analysis=show)
		periods = [frappe._dict(key=f"period_{i}", label=str(i), from_date=date(2026, m, 1), to_date=date(2026, m, 28)) for i, m in enumerate((1, 2))]
		with ExitStack() as stack:
			for name, value in (("normalize", filters), ("get_scope", (projects, accounts, {a.name for a in accounts})),
				("get_periods", periods), ("read_entries", entries), ("get_company_currency", "AED")):
				stack.enter_context(patch.object(report, name, return_value=value))
			stack.enter_context(patch.object(report.frappe, "get_precision", return_value=2))
			return report.execute(filters)

	def test_centers_consolidate_accounts_without_changing_totals(self):
		columns, rows, *_ = self.example()
		self.assertEqual(next(c for c in columns if c["fieldname"] == "label")["label"], "Cost Center")
		centers = [r for r in rows if r.get("_exact_center")]
		self.assertFalse(any(r.get("account") or r.get("indent", 0) > 1 for r in rows))
		self.assertEqual(len({(r["project"], r["cost_center"]) for r in centers}), len(centers))
		block = next(r for r in centers if r["project"] == "P1" and r["cost_center"] == "Block")
		self.assertEqual((block["expenses"], block["period_0"], block["period_1"]), (19497.94, 10000, 9497.94))
		self.assertAlmostEqual(block["expense_contribution"], 2.08042856, places=5)
		self.assertAlmostEqual(block["direct_expense_contribution"], block["expense_contribution"])
		self.assertAlmostEqual(block["expense_project_income"], 1.259201, places=5)
		self.assertEqual(block["income_contribution"], 0)
		self.assertIsNone(block["margin"])
		for parent in [r for r in rows if r.get("_project_summary")]:
			children = [r for r in centers if r["project"] == parent["project"]]
			for field in (*report.AMOUNTS, "expenses", "period_0", "period_1", "total"):
				self.assertAlmostEqual(sum(r[field] for r in children), parent[field])
			self.assertAlmostEqual(sum(r["expense_contribution"] for r in children), 100)
		parent = next(r for r in rows if r.get("_project_summary") and r["project"] == "P1")
		self.assertEqual(parent["gross_profit"], 611229.32)
		self.assertAlmostEqual(parent["gross_margin"], 39.4739347, places=4)
		self.assertEqual(parent["gross_margin"], parent["margin"])
		grand = rows[-1]
		self.assertEqual(grand["expenses"], 937220.24)
		self.assertEqual(grand["margin"], grand["net_profit"] / grand["income"] * 100)

	def test_missing_center_zero_denominators_and_reversals(self):
		_, rows, *_ = self.example(entries=[("P1", None, "D1", 100, 1), ("P1", "Refund", "D2", -20, 2)])
		centers = [r for r in rows if r.get("_exact_center")]
		self.assertEqual(centers[0]["label"], "No Cost Center")
		self.assertEqual([r["expense_contribution"] for r in centers], [125, -25])
		for row in centers:
			self.assertIsNone(row["income_contribution"])
			self.assertIsNone(row["expense_project_income"])
			self.assertIsNone(row["margin"])
		_, rows, *_ = self.example(entries=[("P1", "Income", "I", -100, 1)])
		self.assertIsNone(rows[1]["expense_contribution"])
		self.assertIsNone(rows[1]["direct_expense_contribution"])

	def test_account_mode_and_optional_percentage_columns(self):
		columns, rows, *_ = self.example("Detailed Accounts", False)
		self.assertTrue(any(r.get("account") == "D1" and r["indent"] == 2 for r in rows))
		self.assertNotIn("expense_contribution", [c["fieldname"] for c in columns])
		self.assertEqual(sum(c["fieldname"] == "margin" for c in columns), 1)
		columns, rows, *_ = self.example(show=False)
		self.assertFalse(any(r.get("account") for r in rows))
		self.assertNotIn("expense_contribution", [c["fieldname"] for c in columns])

	def test_equal_postings_are_not_deduplicated_and_zero_net_cost_is_undefined(self):
		_, rows, *_ = self.example(entries=[("P1", "Block", "D1", 20, 1), ("P1", "Block", "D2", 20, 1), ("P1", "Credit", "D1", -40, 2)])
		block = next(r for r in rows if r.get("cost_center") == "Block")
		self.assertEqual(block["expenses"], 40)
		self.assertIsNone(block["expense_contribution"])
		self.assertIsNone(block["direct_expense_contribution"])
		self.assertEqual(rows[-1]["expenses"], 0)

	def test_native_excel_is_numeric_styled_and_not_scaled(self):
		from frappe.desk.query_report import build_xlsx_data
		from frappe.utils.xlsxutils import make_xlsx
		from openpyxl import load_workbook

		columns, rows, *_ = self.example()
		data = frappe._dict(columns=columns, result=rows, report_name="Project Profitability", filters={}, add_total_row=False)
		values, widths, styles = build_xlsx_data(data, [], ignore_visible_idx=True, build_styles=True)
		workbook = load_workbook(BytesIO(make_xlsx(values, "Profitability", column_widths=widths, styles=styles).getvalue()))
		sheet = workbook.active
		headers = [cell.value for cell in sheet[1]]
		idx = headers.index("Expense Contribution %") + 1
		block_index = next(i + 2 for i, r in enumerate(rows) if r.get("cost_center") == "Block" and r["project"] == "P1")
		cell = sheet.cell(block_index, idx)
		self.assertIsInstance(cell.value, float)
		self.assertAlmostEqual(cell.value, 2.08042856, places=5)
		self.assertEqual(cell.number_format, '0.00"%"')
		self.assertTrue(sheet.cell(2, 1).font.bold)
		margin_index = headers.index("Net Profit Margin %") + 1
		self.assertIsNone(sheet.cell(block_index, margin_index).value)
